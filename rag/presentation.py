import json
import os
import uuid
from pathlib import Path

from dotenv import load_dotenv
from openai import OpenAI
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

load_dotenv()

OUTPUT_DIR = Path("generated/presentations")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# THEMES
# ============================================================

THEMES = {
    "modern": {
        "bg": "0B0B10",
        "surface": "17151F",
        "surface2": "211D2D",
        "accent": "A855F7",
        "accent2": "6366F1",
        "text": "FFFFFF",
        "muted": "B7B2C5",
    },

    "corporate": {
        "bg": "F5F7FA",
        "surface": "FFFFFF",
        "surface2": "E9EEF5",
        "accent": "2563EB",
        "accent2": "0F4C81",
        "text": "172033",
        "muted": "64748B",
    },

    "academic": {
        "bg": "F8F6F0",
        "surface": "FFFFFF",
        "surface2": "EDE8DA",
        "accent": "7C3AED",
        "accent2": "92400E",
        "text": "27221B",
        "muted": "6B6257",
    },

    "minimal": {
        "bg": "FFFFFF",
        "surface": "F5F5F5",
        "surface2": "EAEAEA",
        "accent": "111111",
        "accent2": "555555",
        "text": "111111",
        "muted": "666666",
    },

    "dark": {
        "bg": "080B12",
        "surface": "111827",
        "surface2": "1F2937",
        "accent": "22D3EE",
        "accent2": "3B82F6",
        "text": "F9FAFB",
        "muted": "9CA3AF",
    },
}


# ============================================================
# HELPERS
# ============================================================

def rgb(hex_value):
    hex_value = hex_value.replace("#", "")

    return RGBColor(
        int(hex_value[0:2], 16),
        int(hex_value[2:4], 16),
        int(hex_value[4:6], 16),
    )


def add_rect(
    slide,
    x,
    y,
    w,
    h,
    fill,
    radius=False,
    line=None,
):
    shape_type = (
        MSO_SHAPE.ROUNDED_RECTANGLE
        if radius
        else MSO_SHAPE.RECTANGLE
    )

    shape = slide.shapes.add_shape(
        shape_type,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)

    if line:
        shape.line.color.rgb = rgb(line)
    else:
        shape.line.fill.background()

    return shape


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=20,
    color="FFFFFF",
    bold=False,
    align=PP_ALIGN.LEFT,
    font="Aptos",
):
    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )

    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p = tf.paragraphs[0]
    p.text = str(text)
    p.alignment = align

    run = p.runs[0]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)

    return box


def add_background(slide, theme):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(theme["bg"])


def add_top_accent(slide, theme):
    add_rect(
        slide,
        0,
        0,
        13.333,
        0.07,
        theme["accent"],
    )


def add_footer(slide, number, theme):
    add_text(
        slide,
        str(number),
        12.15,
        7.02,
        0.55,
        0.25,
        size=9,
        color=theme["muted"],
        align=PP_ALIGN.RIGHT,
    )


def clean_text(text):
    text = str(text).strip()

    prefixes = [
        "Point 1:",
        "Point 2:",
        "Point 3:",
        "Point 4:",
        "Point 5:",
        "Point 6:",
        "• ",
        "- ",
        "* ",
    ]

    for prefix in prefixes:
        if text.startswith(prefix):
            text = text[len(prefix):].strip()

    return text


# ============================================================
# TITLE SLIDE
# ============================================================

def create_title_slide(
    prs,
    slide_data,
    topic,
    theme,
    style,
):
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_background(slide, theme)

    # Decorative visual blocks
    add_rect(
        slide,
        9.65,
        -0.5,
        4.4,
        4.4,
        theme["accent"],
        radius=True,
    )

    add_rect(
        slide,
        10.7,
        4.15,
        3.1,
        3.1,
        theme["accent2"],
        radius=True,
    )

    # Small accent line
    add_rect(
        slide,
        0.8,
        1.15,
        1.0,
        0.07,
        theme["accent"],
    )

    title = slide_data.get(
        "title",
        topic,
    )

    add_text(
        slide,
        title,
        0.8,
        1.55,
        8.2,
        1.8,
        size=38,
        color=theme["text"],
        bold=True,
    )

    add_text(
        slide,
        topic,
        0.85,
        3.45,
        7.5,
        0.75,
        size=18,
        color=theme["muted"],
    )

    add_text(
        slide,
        "AI-GENERATED PRESENTATION",
        0.85,
        5.85,
        4.0,
        0.35,
        size=10,
        color=theme["accent"],
        bold=True,
    )

    return slide


# ============================================================
# STANDARD CONTENT SLIDE
# ============================================================

def create_content_slide(
    prs,
    slide_data,
    slide_number,
    theme,
    layout="cards",
):
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_background(slide, theme)
    add_top_accent(slide, theme)

    title = slide_data.get("title", "")

    add_text(
        slide,
        title,
        0.72,
        0.42,
        11.8,
        0.75,
        size=30,
        color=theme["text"],
        bold=True,
    )

    # Small title underline
    add_rect(
        slide,
        0.75,
        1.25,
        0.75,
        0.055,
        theme["accent"],
    )

    content = slide_data.get(
        "content",
        [],
    )

    content = [
        clean_text(item)
        for item in content
        if str(item).strip()
    ]

    if layout == "timeline":
        _timeline_layout(
            slide,
            content,
            theme,
        )

    elif layout == "comparison":
        _comparison_layout(
            slide,
            content,
            theme,
        )

    elif layout == "kpi":
        _kpi_layout(
            slide,
            content,
            theme,
        )

    elif layout == "columns":
        _columns_layout(
            slide,
            content,
            theme,
        )

    elif layout == "highlight":
        _highlight_layout(
            slide,
            content,
            theme,
        )

    else:
        _card_layout(
            slide,
            content,
            theme,
        )

    add_footer(
        slide,
        slide_number,
        theme,
    )

    return slide


# ============================================================
# CARD LAYOUT
# ============================================================

def _card_layout(
    slide,
    content,
    theme,
):
    content = content[:6]

    count = len(content)

    if count <= 3:
        cols = max(count, 1)
        rows = 1
        card_h = 3.8

    elif count == 4:
        cols = 2
        rows = 2
        card_h = 2.25

    else:
        cols = 3
        rows = 2
        card_h = 2.25

    gap_x = 0.28
    gap_y = 0.28
    margin_x = 0.75

    available_w = 11.8

    card_w = (
        available_w
        - gap_x * (cols - 1)
    ) / cols

    start_y = 1.62

    for i, item in enumerate(content):

        row = i // cols
        col = i % cols

        x = (
            margin_x
            + col * (card_w + gap_x)
        )

        y = (
            start_y
            + row * (card_h + gap_y)
        )

        accent = (
            theme["accent"]
            if i % 2 == 0
            else theme["accent2"]
        )

        # Card
        add_rect(
            slide,
            x,
            y,
            card_w,
            card_h,
            theme["surface"],
            radius=True,
        )

        # Accent strip
        add_rect(
            slide,
            x,
            y,
            0.075,
            card_h,
            accent,
        )

        # Number
        add_text(
            slide,
            f"{i + 1:02d}",
            x + 0.3,
            y + 0.22,
            0.7,
            0.35,
            size=11,
            color=accent,
            bold=True,
        )

        # Content
        add_text(
            slide,
            item,
            x + 0.3,
            y + 0.75,
            card_w - 0.6,
            card_h - 0.95,
            size=18 if count <= 4 else 15,
            color=theme["text"],
        )


# ============================================================
# TWO COLUMN LAYOUT
# ============================================================

def _columns_layout(
    slide,
    content,
    theme,
):
    content = content[:4]

    if not content:
        return

    for i, item in enumerate(content):

        col = i % 2
        row = i // 2

        x = (
            0.8
            if col == 0
            else 6.78
        )

        y = 1.7 + row * 2.35

        accent = (
            theme["accent"]
            if col == 0
            else theme["accent2"]
        )

        add_text(
            slide,
            f"{i + 1:02d}",
            x,
            y,
            0.65,
            0.35,
            size=11,
            color=accent,
            bold=True,
        )

        add_text(
            slide,
            item,
            x,
            y + 0.52,
            5.35,
            1.35,
            size=18,
            color=theme["text"],
        )

        add_rect(
            slide,
            x,
            y + 1.98,
            5.2,
            0.035,
            theme["surface2"],
        )


# ============================================================
# HIGHLIGHT LAYOUT
# ============================================================

def _highlight_layout(
    slide,
    content,
    theme,
):
    content = content[:4]

    if not content:
        return

    # Main statement
    add_rect(
        slide,
        0.8,
        1.65,
        11.7,
        1.65,
        theme["surface"],
        radius=True,
    )

    add_rect(
        slide,
        0.8,
        1.65,
        0.08,
        1.65,
        theme["accent"],
    )

    add_text(
        slide,
        content[0],
        1.2,
        1.85,
        10.7,
        1.15,
        size=24,
        color=theme["text"],
        bold=True,
    )

    # Supporting points
    remaining = content[1:]

    for i, item in enumerate(remaining):

        x = 0.85 + i * 4.0

        add_text(
            slide,
            f"{i + 1:02d}",
            x,
            3.85,
            0.6,
            0.35,
            size=11,
            color=theme["accent"],
            bold=True,
        )

        add_text(
            slide,
            item,
            x,
            4.35,
            3.45,
            1.35,
            size=16,
            color=theme["text"],
        )


# ============================================================
# TIMELINE
# ============================================================

def _timeline_layout(
    slide,
    content,
    theme,
):
    content = content[:6]

    if not content:
        return

    line_x = 1.15
    start_y = 1.7
    spacing = min(
        0.82,
        4.5 / max(len(content), 1),
    )

    add_rect(
        slide,
        line_x,
        start_y + 0.25,
        0.05,
        spacing * max(
            len(content) - 1,
            1,
        ) + 0.35,
        theme["accent"],
    )

    for i, item in enumerate(content):

        y = start_y + i * spacing

        add_rect(
            slide,
            0.9,
            y,
            0.55,
            0.55,
            theme["accent"],
            radius=True,
        )

        add_text(
            slide,
            str(i + 1),
            0.9,
            y,
            0.55,
            0.55,
            size=11,
            color="FFFFFF",
            bold=True,
            align=PP_ALIGN.CENTER,
        )

        add_text(
            slide,
            item,
            1.75,
            y - 0.03,
            10.25,
            0.7,
            size=17,
            color=theme["text"],
        )


# ============================================================
# COMPARISON
# ============================================================

def _comparison_layout(
    slide,
    content,
    theme,
):
    content = content[:6]

    if not content:
        return

    midpoint = max(
        1,
        len(content) // 2,
    )

    left = content[:midpoint]
    right = content[midpoint:]

    columns = [
        (
            0.75,
            left,
            theme["accent"],
            "OPTION A",
        ),
        (
            6.75,
            right,
            theme["accent2"],
            "OPTION B",
        ),
    ]

    for x, items, accent, label in columns:

        add_rect(
            slide,
            x,
            1.65,
            5.75,
            4.85,
            theme["surface"],
            radius=True,
        )

        add_rect(
            slide,
            x,
            1.65,
            5.75,
            0.08,
            accent,
        )

        add_text(
            slide,
            label,
            x + 0.35,
            1.95,
            4.7,
            0.4,
            size=11,
            color=accent,
            bold=True,
        )

        for i, item in enumerate(items):

            y = 2.55 + i * 0.92

            add_text(
                slide,
                "✓",
                x + 0.35,
                y,
                0.4,
                0.4,
                size=15,
                color=accent,
                bold=True,
            )

            add_text(
                slide,
                item,
                x + 0.85,
                y - 0.02,
                4.45,
                0.7,
                size=15,
                color=theme["text"],
            )


# ============================================================
# KPI / STATISTICS
# ============================================================

def _kpi_layout(
    slide,
    content,
    theme,
):
    content = content[:4]

    if not content:
        return

    count = len(content)

    gap = 0.28
    card_w = (
        11.8
        - gap * (count - 1)
    ) / count

    start_x = 0.75

    for i, item in enumerate(content):

        x = start_x + i * (
            card_w + gap
        )

        accent = (
            theme["accent"]
            if i % 2 == 0
            else theme["accent2"]
        )

        add_rect(
            slide,
            x,
            1.7,
            card_w,
            4.4,
            theme["surface"],
            radius=True,
        )

        add_rect(
            slide,
            x,
            1.7,
            card_w,
            0.09,
            accent,
        )

        add_text(
            slide,
            f"0{i + 1}",
            x + 0.3,
            2.05,
            0.8,
            0.35,
            size=11,
            color=accent,
            bold=True,
        )

        add_text(
            slide,
            item,
            x + 0.3,
            2.7,
            card_w - 0.6,
            2.6,
            size=18,
            color=theme["text"],
            bold=True,
        )


# ============================================================
# SMART LAYOUT SELECTION
# ============================================================

def choose_layout(
    slide_data,
    slide_index=1,
):
    title = str(
        slide_data.get(
            "title",
            "",
        )
    ).lower()

    content = slide_data.get(
        "content",
        [],
    )

    text = " ".join(
        [title]
        + [str(x) for x in content]
    ).lower()

    # Timeline / process
    if any(
        word in text
        for word in [
            "step",
            "process",
            "phase",
            "workflow",
            "evolution",
            "timeline",
            "stages",
            "history",
        ]
    ):
        return "timeline"

    # Comparison
    if any(
        word in text
        for word in [
            "comparison",
            " vs ",
            "versus",
            "difference",
            "advantages",
            "disadvantages",
            "pros and cons",
        ]
    ):
        return "comparison"

    # Statistics
    if any(
        word in text
        for word in [
            "statistics",
            "metrics",
            "key numbers",
            "kpi",
            "market size",
            "percentage",
            "growth rate",
            "data points",
        ]
    ):
        return "kpi"

    # Application / examples
    if any(
        word in text
        for word in [
            "applications",
            "use cases",
            "industries",
            "examples",
            "real-world",
        ]
    ):
        return "columns"

    # Conclusion / takeaway
    if any(
        word in text
        for word in [
            "conclusion",
            "key takeaway",
            "summary",
            "final thoughts",
            "future outlook",
        ]
    ):
        return "highlight"

    return "cards"


# ============================================================
# LLM CONTENT GENERATION
# ============================================================

def generate_slide_content(
    topic,
    context="",
    slide_count=6,
    audience="business",
    style="modern",
    tone="professional",
):
    prompt = f"""
Create a high-quality professional PowerPoint presentation.

TOPIC:
{topic}

AUDIENCE:
{audience}

STYLE:
{style}

TONE:
{tone}

NUMBER OF SLIDES:
{slide_count}

{f'''
SOURCE DOCUMENT CONTEXT:
{context}

IMPORTANT:
Use ONLY the information contained in the source document when answering.
Do not invent facts, statistics, dates, names, or claims.
''' if context else ""}

Return ONLY valid JSON in this exact structure:

{{
    "slides": [
        {{
            "title": "Clear slide title",
            "content": [
                "Concise point",
                "Concise point",
                "Concise point"
            ]
        }}
    ]
}}

CONTENT RULES:

- Generate exactly {slide_count} slides.
- Slide 1 must be a strong introduction/title slide.
- Every other slide must communicate ONE clear idea.
- Use 3-5 points per content slide.
- Keep each point between roughly 8 and 18 words.
- Never write paragraph-style bullets.
- Avoid generic filler such as "In today's world".
- Avoid repeating the same information on multiple slides.
- Prefer concrete explanations over vague statements.
- Include useful examples when appropriate.
- Include measurable facts only when they are actually supported.
- Use specific terminology relevant to the topic.
- Make titles descriptive and meaningful.
- Do not use "Point 1", "Point 2", etc.
- Do not use markdown.
- Do not use emojis.
- Do not use bullet symbols.
- Do not use code fences.
- Do not include speaker notes.
- Do not include explanations outside the JSON.

STRUCTURE RULES:

Use a logical narrative.

For a general topic, prefer a structure such as:

1. Introduction
2. Background / Evolution
3. Core concepts
4. Technologies / Components
5. Applications / Impact
6. Challenges / Future / Conclusion

Adapt this structure intelligently to the topic rather than copying it blindly.

Make the content presentation-ready and suitable for a real academic or professional presentation.
"""

    # --------------------------------------------------------
    # GROQ
    # --------------------------------------------------------

    groq_key = os.getenv(
        "GROQ_API_KEY"
    )

    if groq_key:

        try:
            print(
                "\nTrying Groq for presentation generation..."
            )

            client = OpenAI(
                api_key=groq_key,
                base_url="https://api.groq.com/openai/v1",
            )

            response = client.chat.completions.create(
                model="openai/gpt-oss-120b",
                messages=[
                    {
                        "role": "user",
                        "content": prompt,
                    }
                ],
                temperature=0.45,
            )

            content = (
                response
                .choices[0]
                .message
                .content
            )

            print("Groq selected.")

            return _parse_slide_json(
                content
            )

        except Exception as e:
            print(
                f"Groq unavailable: {e}"
            )

    # --------------------------------------------------------
    # OPENROUTER
    # --------------------------------------------------------

    openrouter_key = os.getenv(
        "OPENROUTER_API_KEY"
    )

    if openrouter_key:

        try:
            print(
                "\nTrying OpenRouter for presentation generation..."
            )

            client = OpenAI(
                api_key=openrouter_key,
                base_url="https://openrouter.ai/api/v1",
            )

            response = client.chat.completions.create(
                model="openrouter/free",
                messages=[
                    {
                        "role": "user",
                        "content": prompt,
                    }
                ],
                temperature=0.45,
            )

            content = (
                response
                .choices[0]
                .message
                .content
            )

            print("OpenRouter selected.")

            return _parse_slide_json(
                content
            )

        except Exception as e:
            print(
                f"OpenRouter unavailable: {e}"
            )

    raise RuntimeError(
        "Unable to generate presentation. "
        "All configured LLM providers are unavailable."
    )


# ============================================================
# JSON PARSER
# ============================================================

def _parse_slide_json(content):

    content = content.strip()

    if content.startswith("```"):

        content = content.replace(
            "```json",
            "",
        )

        content = content.replace(
            "```",
            "",
        )

        content = content.strip()

    data = json.loads(content)

    if "slides" not in data:
        raise ValueError(
            "LLM response does not contain 'slides'."
        )

    if not isinstance(
        data["slides"],
        list,
    ):
        raise ValueError(
            "'slides' must be a list."
        )

    if not data["slides"]:
        raise ValueError(
            "LLM returned an empty slide list."
        )

    return data["slides"]


# ============================================================
# PRESENTATION CREATOR
# ============================================================

def create_presentation(
    topic,
    slides,
    include_notes=False,
    style="modern",
    audience="business",
    include_images=False,
    include_sources=False,
    sources=None,
):
    """
    Create a polished PowerPoint presentation.

    Images are intentionally disabled.
    The renderer focuses on:
    - typography
    - spacing
    - cards
    - timelines
    - comparisons
    - statistics
    - professional layouts
    """

    presentation_id = uuid.uuid4().hex

    filename = (
        f"{presentation_id}.pptx"
    )

    output_path = (
        OUTPUT_DIR / filename
    )

    prs = Presentation()

    # Widescreen 16:9
    prs.slide_width = Inches(
        13.333
    )

    prs.slide_height = Inches(
        7.5
    )

    theme = THEMES.get(
        str(style).lower(),
        THEMES["modern"],
    )

    # --------------------------------------------------------
    # CREATE SLIDES
    # --------------------------------------------------------

    for index, slide_data in enumerate(
        slides
    ):

        if index == 0:

            create_title_slide(
                prs,
                slide_data,
                topic,
                theme,
                style,
            )

        else:

            layout = choose_layout(
                slide_data,
                index,
            )

            create_content_slide(
                prs,
                slide_data,
                index + 1,
                theme,
                layout,
            )

    # --------------------------------------------------------
    # SOURCES SLIDE
    # --------------------------------------------------------

    if include_sources and sources:

        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        add_background(
            slide,
            theme,
        )

        add_top_accent(
            slide,
            theme,
        )

        add_text(
            slide,
            "Sources",
            0.72,
            0.42,
            11.8,
            0.75,
            size=30,
            color=theme["text"],
            bold=True,
        )

        add_rect(
            slide,
            0.75,
            1.25,
            0.75,
            0.055,
            theme["accent"],
        )

        for i, source in enumerate(
            sources
        ):

            y = 1.7 + i * 0.62

            add_text(
                slide,
                f"[{i + 1}]",
                0.85,
                y,
                0.55,
                0.4,
                size=11,
                color=theme["accent"],
                bold=True,
            )

            add_text(
                slide,
                str(source),
                1.45,
                y,
                10.7,
                0.5,
                size=14,
                color=theme["muted"],
            )

        add_footer(
            slide,
            len(prs.slides),
            theme,
        )

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    prs.save(
        output_path
    )

    return {
        "presentation_id": presentation_id,
        "filename": filename,
        "path": str(output_path),
        "slide_count": len(
            prs.slides
        ),
    }